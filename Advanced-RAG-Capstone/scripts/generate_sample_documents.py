"""
Create the synthetic TechNova sample documents (PDF and DOCX).

The project ships with these files already generated, so you normally never
need to run this script. It exists so you can see exactly where the sample
content came from, and regenerate it if you delete or edit the files.

Run it with:
    python scripts/generate_sample_documents.py

The content is invented. Its only job is to give the RAG pipeline realistic
company documents to search through, in three different file formats.
"""

import sys
from pathlib import Path
from typing import List, Tuple

# Allow "import config" when running this file directly from the scripts folder.
sys.path.insert(0, str(Path(__file__).parent.parent))

import config  # noqa: E402

# Each document is a list of (heading, paragraphs) pairs. Headings matter:
# `src/ingestion.guess_section` picks them up and turns them into the
# "Section: ..." part of a citation.
Section = Tuple[str, List[str]]


EMPLOYEE_HANDBOOK: List[Section] = [
    (
        "Welcome to TechNova",
        [
            "TechNova Solutions is a software company founded in 2016 in Austin, "
            "Texas. We build NovaCloud, a platform that helps mid-sized companies "
            "run and monitor their data pipelines. As of 2024 we employ 240 people "
            "across three offices: Austin, Dublin and Singapore.",
            "This handbook describes how we work together. It applies to all "
            "permanent employees. Contractors are covered by the terms of their "
            "individual agreements.",
        ],
    ),
    (
        "Working Hours",
        [
            "TechNova operates on flexible hours. Employees are expected to be "
            "available during the core hours of 10:00 to 16:00 in their local time "
            "zone, and to work an average of 40 hours per week.",
            "Teams hold a daily stand-up inside core hours. Meetings outside core "
            "hours should be recorded so colleagues in other regions can catch up.",
        ],
    ),
    (
        "Remote Work",
        [
            "All engineering roles are remote-friendly. Employees may work remotely "
            "up to four days per week, and fully remote arrangements are approved "
            "case by case by the reporting manager and the People team.",
            "Employees working remotely from another country for more than 30 days "
            "in a calendar year must notify the People team first, because this can "
            "create tax obligations for both the employee and TechNova.",
        ],
    ),
    (
        "Annual Leave",
        [
            "Every full-time employee receives 24 days of paid annual leave per "
            "calendar year, in addition to local public holidays. Leave is accrued "
            "monthly at two days per month.",
            "Up to five unused leave days may be carried over into the next year "
            "and must be used before 31 March. Leave requests are submitted in the "
            "People portal and are normally approved within three working days.",
        ],
    ),
    (
        "Sick Leave",
        [
            "Employees receive up to 10 days of paid sick leave per year. A medical "
            "certificate is required for any absence longer than three consecutive "
            "working days.",
            "Mental health days count as sick leave. Employees do not have to "
            "explain the nature of an illness to their manager.",
        ],
    ),
    (
        "Parental Leave",
        [
            "TechNova offers 16 weeks of fully paid parental leave to the primary "
            "caregiver and 6 weeks to the secondary caregiver. Leave can be taken "
            "in up to three separate blocks during the first 18 months.",
            "Employees returning from parental leave may work a reduced four-day "
            "week at 80 percent pay for the first three months.",
        ],
    ),
    (
        "Learning Budget",
        [
            "Each employee has an annual learning budget of 1500 USD. It can be "
            "spent on courses, conference tickets, books and certifications. The "
            "budget does not carry over to the following year.",
            "Purchases above 500 USD need manager approval in advance. Submit "
            "receipts through the Expenses portal within 30 days of purchase.",
        ],
    ),
    (
        "Expense Reimbursement",
        [
            "Approved business expenses are reimbursed with the next monthly "
            "payroll run, provided the claim is submitted before the 25th of the "
            "month. Claims submitted later are paid the following month.",
            "Travel must be booked through the TechNova travel portal. Economy "
            "class is standard for flights under six hours; business class may be "
            "booked for longer flights with director approval.",
        ],
    ),
    (
        "Equipment Policy",
        [
            "New employees receive a laptop, an external monitor, a keyboard and a "
            "mouse. Engineers may choose between a MacBook Pro and a Dell XPS "
            "running Ubuntu. Laptops are refreshed every three years.",
            "Employees also receive a one-time home office allowance of 600 USD for "
            "a desk and chair. Equipment remains the property of TechNova and must "
            "be returned when employment ends.",
        ],
    ),
    (
        "Performance Reviews",
        [
            "Performance reviews happen twice a year, in March and September. Each "
            "review combines a self-assessment, manager feedback and peer feedback "
            "from at least two colleagues.",
            "Promotion decisions and salary adjustments are made during the March "
            "cycle. Compensation changes take effect from 1 May.",
        ],
    ),
    (
        "Code of Conduct",
        [
            "TechNova does not tolerate harassment, discrimination or bullying of "
            "any kind. Concerns can be raised with a manager, with the People team, "
            "or anonymously through the ethics hotline.",
            "All reports are investigated within 10 working days. Retaliation "
            "against anyone who raises a concern in good faith is itself a serious "
            "breach of this policy.",
        ],
    ),
]


PRODUCT_GUIDE: List[Section] = [
    (
        "NovaCloud Overview",
        [
            "NovaCloud is TechNova's data pipeline platform. It lets teams build, "
            "schedule and monitor pipelines that move data between databases, "
            "object storage and SaaS applications, without managing servers.",
            "The platform has three parts: the Pipeline Builder, the Scheduler and "
            "the Observability Hub. The current release is NovaCloud 4.2, shipped "
            "in February 2024.",
        ],
    ),
    (
        "Pipeline Builder",
        [
            "The Pipeline Builder is a visual editor where a pipeline is assembled "
            "from connectors, transforms and destinations. Every pipeline can also "
            "be exported as a YAML file and stored in Git.",
            "Transforms are written in SQL or Python. Python transforms run in an "
            "isolated sandbox with a 10 minute execution limit and 4 GB of memory.",
        ],
    ),
    (
        "Scheduler",
        [
            "Pipelines can run on a fixed schedule using cron expressions, on an "
            "event trigger such as a new file landing in object storage, or on "
            "demand through the REST API.",
            "Failed runs are retried three times with exponential backoff. After "
            "the third failure the pipeline is marked as failed and an alert is "
            "sent to the configured notification channel.",
        ],
    ),
    (
        "Observability Hub",
        [
            "The Observability Hub shows run history, row counts, execution time "
            "and data quality checks for every pipeline. Metrics are retained for "
            "13 months on all paid plans.",
            "Data quality checks include null checks, uniqueness checks, range "
            "checks and freshness checks. A failing check can either warn or stop "
            "the pipeline, depending on how it is configured.",
        ],
    ),
    (
        "Connectors",
        [
            "NovaCloud 4.2 ships with 68 connectors, including PostgreSQL, MySQL, "
            "Snowflake, BigQuery, Amazon S3, Google Cloud Storage, Salesforce, "
            "HubSpot, Stripe and Zendesk.",
            "Customers can build private connectors with the Connector SDK. A "
            "private connector is visible only inside the customer's workspace and "
            "is not reviewed by TechNova.",
        ],
    ),
    (
        "Pricing Tiers",
        [
            "NovaCloud has three subscription tiers. Starter costs 99 USD per month "
            "and includes 3 pipelines, 1 workspace and 10 GB of monthly data "
            "transfer.",
            "Growth costs 499 USD per month and includes 25 pipelines, 3 workspaces "
            "and 250 GB of monthly data transfer. Enterprise is priced per contract "
            "and includes unlimited pipelines, unlimited workspaces, single sign-on "
            "and a private deployment option.",
            "Data transfer above the included allowance is billed at 0.09 USD per "
            "GB. Usage is measured per calendar month and shown live in the Billing "
            "page.",
        ],
    ),
    (
        "Platform Limits",
        [
            "A single pipeline may contain up to 150 steps. The maximum size of one "
            "file processed in a single step is 5 GB. The REST API is rate limited "
            "to 600 requests per minute per workspace.",
            "Limits can be raised for Enterprise customers by contacting the "
            "Technical Account Manager.",
        ],
    ),
    (
        "Integrations and Alerting",
        [
            "Alerts can be delivered to email, Slack, Microsoft Teams, PagerDuty "
            "and any HTTPS webhook. Each pipeline can have its own alerting rules.",
            "NovaCloud also exposes a Prometheus-compatible metrics endpoint so "
            "customers can pull pipeline metrics into their own monitoring stack.",
        ],
    ),
    (
        "Getting Started",
        [
            "After signing up, create a workspace, connect a source, choose a "
            "destination and run the pipeline once manually. The guided tutorial in "
            "the console walks through this in about 15 minutes.",
            "Free trials last 14 days and include all Growth tier features. No "
            "credit card is required to start a trial.",
        ],
    ),
    (
        "Deprecation and Versioning",
        [
            "The NovaCloud REST API is versioned. Version v1 is deprecated and will "
            "be switched off on 31 December 2024. Version v2 is the current, "
            "supported version.",
            "TechNova gives at least 12 months notice before removing an API "
            "version, and 6 months notice before removing a connector.",
        ],
    ),
]


SECURITY_POLICY: List[Section] = [
    (
        "Purpose and Scope",
        [
            "This Information Security Policy applies to all TechNova employees, "
            "contractors and systems that store or process customer data. It is "
            "reviewed every 12 months by the Security team and approved by the CTO.",
            "Breaking this policy may lead to disciplinary action, up to and "
            "including termination of employment or contract.",
        ],
    ),
    (
        "Password Policy",
        [
            "Passwords must be at least 14 characters long and must not be reused "
            "across systems. All TechNova employees must store work passwords in "
            "the company password manager, 1Password.",
            "Passwords are not rotated on a fixed schedule. Following current NIST "
            "guidance, TechNova only forces a password change when there is "
            "evidence of compromise. Sharing a password with another person, "
            "including a colleague, is never permitted.",
        ],
    ),
    (
        "Multi-Factor Authentication",
        [
            "Multi-factor authentication is mandatory for every TechNova system "
            "that supports it, including email, the production cloud console, the "
            "code repository and the VPN.",
            "Hardware security keys or an authenticator app must be used. SMS codes "
            "are not accepted as a second factor because of the risk of SIM "
            "swapping attacks.",
        ],
    ),
    (
        "Access Control",
        [
            "Access to production systems follows the principle of least privilege. "
            "Engineers receive read-only access by default; write access to "
            "production is granted for a maximum of eight hours through the "
            "just-in-time access tool.",
            "All access rights are reviewed quarterly. Accounts belonging to people "
            "who leave TechNova are disabled within two hours of their final "
            "working day.",
        ],
    ),
    (
        "Data Encryption",
        [
            "Customer data is encrypted in transit with TLS 1.2 or higher, and at "
            "rest with AES-256. Encryption keys are managed in the cloud provider's "
            "key management service and rotated annually.",
            "Customer data is never copied to a laptop or a personal device. "
            "Debugging is done through approved tooling that masks personally "
            "identifiable information.",
        ],
    ),
    (
        "Device Security",
        [
            "All company laptops have full disk encryption, an automatic screen "
            "lock after five minutes, and the company endpoint protection agent "
            "installed. Employees must install operating system security updates "
            "within 14 days of release.",
            "Lost or stolen devices must be reported to security@technova.example "
            "immediately so the device can be remotely wiped.",
        ],
    ),
    (
        "Incident Response",
        [
            "A security incident is any event that may compromise the "
            "confidentiality, integrity or availability of customer data. Anyone "
            "who suspects an incident must report it within one hour.",
            "The Security team triages the report, assigns a severity and leads the "
            "response. For confirmed breaches of personal data, affected customers "
            "and the relevant supervisory authority are notified within 72 hours, "
            "as required by GDPR.",
        ],
    ),
    (
        "Compliance and Certification",
        [
            "TechNova holds a SOC 2 Type II report, renewed annually, and is "
            "certified against ISO 27001. Customers can request both reports under "
            "a non-disclosure agreement from their account manager.",
            "Penetration tests are carried out by an external firm twice a year. A "
            "summary of findings is available to Enterprise customers on request.",
        ],
    ),
    (
        "Vendor Security",
        [
            "Any new vendor that will process customer data must pass a security "
            "review before a contract is signed. The review covers certifications, "
            "data location, sub-processors and the vendor's own incident history.",
            "A register of all approved sub-processors is published on the TechNova "
            "trust page and updated within 30 days of any change.",
        ],
    ),
    (
        "Security Training",
        [
            "Every employee completes security awareness training during onboarding "
            "and again every 12 months. Engineers additionally complete secure "
            "coding training each year.",
            "TechNova runs simulated phishing exercises each quarter. Employees who "
            "click a simulated phishing link are enrolled in a short refresher "
            "course rather than penalised.",
        ],
    ),
]


def write_pdf(sections: List[Section], title: str, output_path: Path) -> None:
    """
    Write a list of sections to a multi-page PDF file.

    We use reportlab's "platypus" layer, which flows paragraphs onto pages
    automatically. That gives us a realistic multi-page PDF, so page numbers
    in the citations are actually meaningful.

    Args:
        sections: (heading, paragraphs) pairs to write.
        title: Title printed on the first page.
        output_path: Where to save the PDF.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(output_path), pagesize=A4, title=title)

    flowables = [Paragraph(title, styles["Title"]), Spacer(1, 18)]
    for heading, paragraphs in sections:
        flowables.append(Paragraph(heading, styles["Heading2"]))
        for text in paragraphs:
            flowables.append(Paragraph(text, styles["BodyText"]))
        flowables.append(Spacer(1, 10))

    document.build(flowables)
    print(f"Created {output_path}")


def write_docx(sections: List[Section], title: str, output_path: Path) -> None:
    """
    Write a list of sections to a Word (.docx) file.

    Having one DOCX in the sample set proves the ingestion step really does
    handle three different formats, not just plain text.

    Args:
        sections: (heading, paragraphs) pairs to write.
        title: Title for the document.
        output_path: Where to save the DOCX.
    """
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading(title, level=0)

    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)

    document.save(str(output_path))
    print(f"Created {output_path}")


def main() -> None:
    """
    Generate every sample document into `data/documents`.

    The TXT sample (technova_support_policy.txt) is already a plain text file
    in the repository, so it does not need generating here.
    """
    config.DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

    write_pdf(
        EMPLOYEE_HANDBOOK,
        "TechNova Employee Handbook 2024",
        config.DOCUMENTS_DIR / "technova_employee_handbook.pdf",
    )
    write_pdf(
        PRODUCT_GUIDE,
        "NovaCloud Product Guide v4.2",
        config.DOCUMENTS_DIR / "technova_product_guide.pdf",
    )
    write_docx(
        SECURITY_POLICY,
        "TechNova Information Security Policy",
        config.DOCUMENTS_DIR / "technova_security_policy.docx",
    )


if __name__ == "__main__":
    main()
